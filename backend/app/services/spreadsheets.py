"""Reading candidate / question spreadsheets (.xlsx, .csv) and writing result reports."""

import csv
import io
import os
import re

from openpyxl import Workbook, load_workbook

from ..errors import ApiError

EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")

CANDIDATE_COLUMNS = {
    "name": ["name", "fullname", "candidatename", "studentname", "applicantname"],
    "first_name": ["firstname", "fname", "givenname"],
    "last_name": ["lastname", "lname", "surname", "familyname"],
    "email": ["email", "emailid", "emailaddress", "mail", "mailid", "candidateemail", "studentemail"],
    "phone": ["phone", "mobile", "phonenumber", "mobilenumber", "contact", "contactnumber", "mobileno", "phoneno"],
}


def _norm(header):
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def read_rows(filename, data):
    """Return (headers, rows) where rows are lists of cell strings."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ApiError("Could not read the Excel file. Save it as .xlsx and try again.")
        ws = wb.worksheets[0]
        all_rows = [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    elif ext == ".csv":
        text = data.decode("utf-8-sig", errors="ignore")
        all_rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]
    else:
        raise ApiError("Upload an .xlsx or .csv file")
    all_rows = [r for r in all_rows if any(c for c in r)]
    if not all_rows:
        raise ApiError("The file is empty")
    headers = all_rows[0]
    rows = [r + [""] * (len(headers) - len(r)) for r in all_rows[1:]]
    return headers, rows


def parse_candidates(filename, data):
    """Parse a candidate sheet. Returns (valid_records, errors)."""
    headers, rows = read_rows(filename, data)
    normed = [_norm(h) for h in headers]
    index = {}
    for field, aliases in CANDIDATE_COLUMNS.items():
        for i, h in enumerate(normed):
            if h in aliases and field not in index:
                index[field] = i
    if "email" not in index:
        raise ApiError(
            "Couldn't find an email column. The first row must contain headers such as Name, Email, Phone.",
            found_headers=headers,
        )
    used = set(index.values())
    extra_cols = [(i, str(h).strip()) for i, h in enumerate(headers) if i not in used and str(h).strip()]

    records, errors, seen = [], [], set()
    for line_no, row in enumerate(rows, start=2):
        email = row[index["email"]].strip().lower()
        if "name" in index:
            name = row[index["name"]].strip()
        else:
            parts = [row[index[k]].strip() for k in ("first_name", "last_name") if k in index]
            name = " ".join(p for p in parts if p)
        if not EMAIL_RE.match(email):
            errors.append({"row": line_no, "error": f"Invalid email '{email}'"})
            continue
        if email in seen:
            errors.append({"row": line_no, "error": f"Duplicate email '{email}' in file"})
            continue
        seen.add(email)
        phone = row[index["phone"]].strip() if "phone" in index else ""
        if phone.endswith(".0") and phone[:-2].isdigit():  # Excel stores numbers as floats
            phone = phone[:-2]
        extra = {h: row[i] for i, h in extra_cols if row[i]}
        records.append({"name": name or email.split("@")[0], "email": email, "phone": phone, "extra": extra})
    return records, errors


def parse_mcqs(filename, data):
    """Parse an MCQ sheet with columns: Question, Option A..F, Answer (letter, number or text), Marks, Topic,
    Difficulty."""
    headers, rows = read_rows(filename, data)
    normed = [_norm(h) for h in headers]

    def col(*names):
        for n in names:
            if n in normed:
                return normed.index(n)
        return None

    q_col = col("question", "questiontext", "questions")
    if q_col is None:
        raise ApiError("Missing 'Question' column", found_headers=headers)
    option_cols = []
    for letter in "abcdef":
        c = col(f"option{letter}", f"opt{letter}", letter, f"choice{letter}")
        if c is not None:
            option_cols.append(c)
    if not option_cols:
        option_cols = [i for i, h in enumerate(normed) if re.fullmatch(r"option\d", h)]
    a_col = col("answer", "correctanswer", "correctoption", "correct", "key")
    if len(option_cols) < 2 or a_col is None:
        raise ApiError("Need at least two option columns (Option A, Option B, ...) and an 'Answer' column")
    m_col, t_col, d_col = col("marks", "mark", "points", "score"), col("topic", "category", "section"), col(
        "difficulty", "level"
    )

    out, errors = [], []
    for line_no, row in enumerate(rows, start=2):
        text = row[q_col].strip()
        options = [row[c].strip() for c in option_cols if row[c].strip()]
        raw = row[a_col].strip()
        answer = None
        if len(raw) == 1 and raw.upper() in "ABCDEF" and "ABCDEF".index(raw.upper()) < len(options):
            answer = "ABCDEF".index(raw.upper())
        elif raw.replace(".0", "").isdigit() and 1 <= int(raw.replace(".0", "")) <= len(options):
            answer = int(raw.replace(".0", "")) - 1
        else:
            lowered = [o.lower() for o in options]
            if raw.lower() in lowered:
                answer = lowered.index(raw.lower())
        if not text or len(options) < 2 or answer is None:
            errors.append({"row": line_no, "error": "Needs a question, 2+ options and a valid answer"})
            continue
        try:
            marks = float(row[m_col]) if m_col is not None and row[m_col] else 1.0
        except ValueError:
            marks = 1.0
        out.append(
            {
                "text": text,
                "options": options,
                "answer": answer,
                "marks": marks,
                "topic": row[t_col].strip() if t_col is not None else "",
                "difficulty": (row[d_col].strip().lower() if d_col is not None else "") or "medium",
            }
        )
    return out, errors


def workbook_bytes(sheets):
    """sheets: list of (title, headers, rows)."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(45, len(str(h)) + 6))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def candidate_template():
    return workbook_bytes(
        [
            (
                "Candidates",
                ["Name", "Email", "Phone", "College", "Graduation Year", "CGPA"],
                [
                    ["Asha Verma", "asha.verma@example.com", "9876543210", "ABC Institute of Technology", 2025, 8.4],
                    ["Rahul Mehta", "rahul.mehta@example.com", "9123456780", "XYZ University", 2024, 7.9],
                ],
            )
        ]
    )


def mcq_template():
    return workbook_bytes(
        [
            (
                "MCQs",
                ["Question", "Option A", "Option B", "Option C", "Option D", "Answer", "Marks", "Topic", "Difficulty"],
                [
                    ["What is 15% of 200?", "15", "30", "25", "20", "B", 1, "Percentages", "easy"],
                    ["Which data structure is FIFO?", "Stack", "Queue", "Tree", "Graph", "B", 1, "CS Basics", "easy"],
                ],
            )
        ]
    )
